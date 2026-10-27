"""Build editable Word reports from validated report objects."""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from datetime import date
from pathlib import Path

from docx import Document as create_document
from docx.document import Document as DocumentObject
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor, Twips
from docx.styles.style import BaseStyle
from docx.table import _Cell, _Row, Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from app.reporting.formatting import format_value
from app.reporting.models import (
    KeyFigure,
    ReportFormat,
    ReportObject,
    ReportValue,
)

DARK_BLUE = "1F3A5F"
SECONDARY_GREY = "666666"
TILE_FILL = "EEF2F7"
TABLE_BORDER = "D0D7E2"
WHITE = "FFFFFF"
BODY_FONT = "Arial"
BODY_SIZE = 10.5


def build_report_document(
    report: ReportObject,
    output_path: Path,
    *,
    chart_path: Path | None = None,
    source_labels: Mapping[str, str] | None = None,
    footer_note: str = "Sample report. All figures are illustrative",
) -> Path:
    """Write a formatted, editable DOCX from an already validated report."""
    output_path = Path(output_path)
    if report.chart is not None and chart_path is None:
        raise ValueError("A chart image path is required when the report has a chart.")
    if chart_path is not None and not Path(chart_path).is_file():
        raise FileNotFoundError(f"Report chart image does not exist: {chart_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)

    document = create_document()
    _configure_page(document)
    _configure_styles(document)
    _add_title_block(document, report.title, report.period_label, report.as_of_date)
    _add_key_figures(document, report.key_figures)
    _add_heading(document, "Summary")
    for paragraph_text in report.summary:
        paragraph = document.add_paragraph(paragraph_text)
        paragraph.paragraph_format.space_after = Pt(6)

    if report.chart is not None and chart_path is not None:
        _add_heading(document, "Revenue trend")
        chart_paragraph = document.add_paragraph()
        chart_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture = chart_paragraph.add_run().add_picture(
            str(chart_path),
            width=Inches(6.2),
        )
        picture._inline.docPr.set("descr", report.chart.title)
        picture._inline.docPr.set("title", report.chart.title)

    for table_data in report.tables:
        _add_heading(document, table_data.title)
        _add_data_table(
            document,
            table_data.columns,
            table_data.formats,
            table_data.rows,
        )

    _add_heading(document, "Findings and recommended actions", new_page=True)
    _add_findings_table(document, report)
    _add_heading(document, "Sources and notes")
    _add_source_and_note_bullets(document, report, source_labels or {})
    _add_page_footer(document, footer_note)

    document.save(str(output_path))
    return output_path


def _configure_page(document: DocumentObject) -> None:
    """Set A4 paper and one-inch margins for each document section."""
    for section in document.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.5)
        section.footer_distance = Inches(0.5)


def _configure_styles(document: DocumentObject) -> None:
    """Set the report's Arial body font and dark-blue heading style."""
    normal = document.styles["Normal"]
    _set_style_font(normal, BODY_FONT)
    _set_style_property(normal, "sz", {"w:val": "21"})

    title = document.styles["Title"]
    _set_style_font(title, BODY_FONT)
    _set_style_property(title, "sz", {"w:val": "52"})
    _set_style_property(title, "b", {})
    _set_style_property(title, "color", {"w:val": DARK_BLUE})

    heading = document.styles["Heading 1"]
    _set_style_font(heading, BODY_FONT)
    _set_style_property(heading, "sz", {"w:val": "30"})
    _set_style_property(heading, "b", {})
    _set_style_property(heading, "color", {"w:val": DARK_BLUE})
    p_pr = heading.element.get_or_add_pPr()
    keep_next = OxmlElement("w:keepNext")
    p_pr.append(keep_next)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "240")
    spacing.set(qn("w:after"), "100")
    p_pr.append(spacing)


def _set_style_font(style: BaseStyle, font_name: str) -> None:
    """Set the Western and East Asian font names on a Word style."""
    style_element = style.element
    r_pr = style_element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:eastAsia"), font_name)


def _set_style_property(
    style: BaseStyle,
    property_name: str,
    attributes: Mapping[str, str],
) -> None:
    """Set one run property in the style XML."""
    r_pr = style.element.get_or_add_rPr()
    property_element = r_pr.find(qn(f"w:{property_name}"))
    if property_element is None:
        property_element = OxmlElement(f"w:{property_name}")
        r_pr.append(property_element)
    for attribute_name, value in attributes.items():
        property_element.set(qn(attribute_name), value)


def _add_title_block(
    document: DocumentObject,
    title: str,
    period_label: str,
    as_of_date: date,
) -> None:
    """Add the title and metadata line with the specified blue rule."""
    title_paragraph = document.add_paragraph(style="Title")
    title_paragraph.paragraph_format.space_after = Pt(2)
    title_paragraph.add_run(title)

    metadata = document.add_paragraph()
    metadata.paragraph_format.space_after = Pt(8)
    metadata.paragraph_format.keep_with_next = True
    metadata_run = metadata.add_run(
        f"{period_label}  |  Prepared by the Business Data Analyst Agent"
        f"  |  Data as of {as_of_date:%d %b %Y}"
    )
    metadata_run.font.name = BODY_FONT
    metadata_run.font.size = Pt(9)
    metadata_run.font.color.rgb = RGBColor.from_string(SECONDARY_GREY)
    _set_paragraph_bottom_border(metadata, DARK_BLUE, "10", "4")


def _add_heading(
    document: DocumentObject,
    text: str,
    *,
    new_page: bool = False,
) -> Paragraph:
    """Add a restylable Heading 1 section heading."""
    paragraph = document.add_paragraph(text, style="Heading 1")
    if new_page:
        paragraph.paragraph_format.page_break_before = True
    return paragraph


def _add_key_figures(document: DocumentObject, figures: Sequence[KeyFigure]) -> None:
    """Render key figures as a single shaded row of editable Word cells."""
    usable_width = _usable_page_width(document)
    tile_table = document.add_table(rows=1, cols=len(figures))
    widths = _weighted_widths(usable_width, len(figures))
    _set_table_widths(tile_table, widths)
    _set_table_borders(tile_table, color=TILE_FILL, size="0")

    for cell, figure, width in zip(
        tile_table.rows[0].cells,
        figures,
        widths,
        strict=True,
    ):
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_shading(cell, TILE_FILL)
        _set_cell_margins(cell, top=110, start=120, bottom=110, end=120)
        label_paragraph = cell.paragraphs[0]
        label_paragraph.paragraph_format.space_after = Pt(3)
        label_run = label_paragraph.add_run(figure.label)
        label_run.font.name = BODY_FONT
        label_run.font.size = Pt(8)
        label_run.font.color.rgb = RGBColor.from_string(SECONDARY_GREY)

        value_paragraph = cell.add_paragraph()
        value_paragraph.paragraph_format.space_after = Pt(0)
        value_run = value_paragraph.add_run(
            format_value(figure.value, figure.format)
        )
        value_run.font.name = BODY_FONT
        value_run.font.size = Pt(17)
        value_run.font.bold = True
        value_run.font.color.rgb = RGBColor.from_string(DARK_BLUE)


def _add_data_table(
    document: DocumentObject,
    columns: Sequence[str],
    formats: Sequence[ReportFormat],
    rows: Sequence[Sequence[ReportValue]],
) -> Table:
    """Add a fixed-width, styled data table with typed cell alignment."""
    table = document.add_table(rows=1, cols=len(columns))
    widths = _weighted_widths(_usable_page_width(document), len(columns))
    _set_table_widths(table, widths)
    _set_table_borders(table, color=TABLE_BORDER, size="4")
    _fill_table_header(table.rows[0].cells, columns)

    for row_values in rows:
        row = table.add_row()
        _prevent_row_split(row)
        for cell, value, value_format, width in zip(
            row.cells,
            row_values,
            formats,
            widths,
            strict=True,
        ):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = format_value(value, value_format)
            _format_cell_paragraph(
                cell.paragraphs[0],
                right_aligned=value_format != "text",
            )
    return table


def _fill_table_header(cells: Sequence[_Cell], columns: Sequence[str]) -> None:
    """Fill a table header with white bold text over a dark-blue shade."""
    for cell, column in zip(cells, columns, strict=True):
        cell.text = column
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_shading(cell, DARK_BLUE)
        _set_cell_margins(cell, top=90, start=90, bottom=90, end=90)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_together = True
            for run in paragraph.runs:
                run.font.name = BODY_FONT
                run.font.size = Pt(8.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor.from_string(WHITE)
    row_element = cells[0]._tc.getparent()
    header_properties = row_element.get_or_add_trPr()
    header_repeat = OxmlElement("w:tblHeader")
    header_repeat.set(qn("w:val"), "true")
    header_properties.append(header_repeat)


def _add_findings_table(document: DocumentObject, report: ReportObject) -> Table:
    """Add a four-column findings table whose rows cannot split across pages."""
    columns = ("Finding", "Evidence", "Recommended action", "Caveat")
    table = document.add_table(rows=1, cols=len(columns))
    widths = _weighted_widths(_usable_page_width(document), len(columns))
    _set_table_widths(table, widths)
    _set_table_borders(table, color=TABLE_BORDER, size="4")
    _fill_table_header(table.rows[0].cells, columns)

    for finding in report.findings:
        row = table.add_row()
        _prevent_row_split(row)
        values = (
            finding.finding,
            finding.evidence,
            finding.action,
            finding.caveat,
        )
        for cell, text, width in zip(row.cells, values, widths, strict=True):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            cell.text = text
            _format_cell_paragraph(cell.paragraphs[0], right_aligned=False)
    return table


def _format_cell_paragraph(
    paragraph: Paragraph,
    *,
    right_aligned: bool,
) -> None:
    """Apply compact Arial typography and requested cell alignment."""
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT if right_aligned else WD_ALIGN_PARAGRAPH.LEFT
    )
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_together = True
    for run in paragraph.runs:
        run.font.name = BODY_FONT
        run.font.size = Pt(8.5)
        _set_run_font(run, BODY_FONT)


def _add_source_and_note_bullets(
    document: DocumentObject,
    report: ReportObject,
    source_labels: Mapping[str, str],
) -> None:
    """List source calls and report notes using Word's editable bullet style."""
    source_refs = {
        figure.source_ref for figure in report.key_figures
    }
    if report.chart is not None:
        source_refs.add(report.chart.source_ref)
    source_refs.update(table.source_ref for table in report.tables)
    for finding in report.findings:
        source_refs.update(finding.source_refs)

    for source_ref in sorted(source_refs):
        source_label = source_labels.get(source_ref, f"Tool result {source_ref}")
        _add_grey_bullet(document, source_label)
    for note in report.notes:
        _add_grey_bullet(document, note)


def _add_grey_bullet(document: DocumentObject, text: str) -> None:
    """Add one source or note as an editable, grey Word bullet paragraph."""
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.font.name = BODY_FONT
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(SECONDARY_GREY)


def _add_page_footer(document: DocumentObject, footer_note: str) -> None:
    """Add the source disclaimer and dynamic page fields to every footer."""
    for section in document.sections:
        footer = section.footer
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(4)
        source_run = paragraph.add_run(f"{footer_note}. Page ")
        source_run.font.name = BODY_FONT
        source_run.font.size = Pt(8)
        source_run.font.color.rgb = RGBColor.from_string(SECONDARY_GREY)
        _append_field(paragraph, "PAGE")
        separator_run = paragraph.add_run(" of ")
        separator_run.font.name = BODY_FONT
        separator_run.font.size = Pt(8)
        separator_run.font.color.rgb = RGBColor.from_string(SECONDARY_GREY)
        _append_field(paragraph, "NUMPAGES")


def _append_field(paragraph: Paragraph, field_name: str) -> None:
    """Append a Word field code that remains dynamic when opened in Word."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = f" {field_name} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    displayed_value = OxmlElement("w:t")
    displayed_value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, displayed_value, end))


def _set_paragraph_bottom_border(
    paragraph: Paragraph,
    color: str,
    size: str,
    space: str,
) -> None:
    """Draw a bottom border under the report metadata line."""
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _set_table_widths(table: Table, widths: Sequence[Twips]) -> None:
    """Set fixed table, grid-column, and cell widths to match the page width."""
    table.autofit = False
    total_twips = sum(int(width.twips) for width in widths)
    tbl_pr = table._tbl.tblPr
    tbl_width = tbl_pr.find(qn("w:tblW"))
    if tbl_width is None:
        tbl_width = OxmlElement("w:tblW")
        tbl_pr.append(tbl_width)
    tbl_width.set(qn("w:w"), str(total_twips))
    tbl_width.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    for grid_column, width in zip(table._tbl.tblGrid.gridCol_lst, widths, strict=True):
        grid_column.set(qn("w:w"), str(width.twips))
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            cell.width = width
            tc_width = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_width is not None:
                tc_width.set(qn("w:w"), str(width.twips))
                tc_width.set(qn("w:type"), "dxa")


def _weighted_widths(total_width: int, column_count: int) -> list[Twips]:
    """Split a table width evenly while assigning the remainder to its last column."""
    width_twips = total_width // 635
    base_width = width_twips // column_count
    widths = [Twips(base_width) for _ in range(column_count)]
    widths[-1] = Twips(width_twips - base_width * (column_count - 1))
    return widths


def _usable_page_width(document: DocumentObject) -> int:
    """Return the page width remaining after left and right margins."""
    section = document.sections[0]
    if (
        section.page_width is None
        or section.left_margin is None
        or section.right_margin is None
    ):
        raise ValueError("Document section is missing its page width or margins.")
    return section.page_width - section.left_margin - section.right_margin


def _set_table_borders(table: Table, *, color: str, size: str) -> None:
    """Set explicit table and cell borders for editable Word tables."""
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single" if size != "0" else "nil")
        border.set(qn("w:sz"), size)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)


def _set_cell_shading(cell: _Cell, fill: str) -> None:
    """Apply a solid fill using the Word cell properties element."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")


def _set_cell_margins(
    cell: _Cell,
    *,
    top: int,
    start: int,
    bottom: int,
    end: int,
) -> None:
    """Set internal cell padding in twentieths of a point."""
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        margin = margins.find(qn(f"w:{side}"))
        if margin is None:
            margin = OxmlElement(f"w:{side}")
            margins.append(margin)
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")


def _prevent_row_split(row: _Row) -> None:
    """Set Word's cantSplit row property to keep table rows intact."""
    row_properties = row._tr.get_or_add_trPr()
    cannot_split = OxmlElement("w:cantSplit")
    row_properties.append(cannot_split)


def _set_run_font(run: Run, font_name: str) -> None:
    """Set explicit East Asian font metadata on a text run."""
    r_pr = run._r.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:eastAsia"), font_name)