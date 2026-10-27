"""Tests for the Step 6 report-generation MCP tool."""

from __future__ import annotations

from pathlib import Path

from docx import Document
import pytest

from app.mcp_servers import business_data_server
from app.state import ToolResult
from tests.test_report_validation import _report, _sources


def _sources_by_call_id() -> dict[str, ToolResult]:
    """Attach workflow tool-call IDs to the evidence fixture."""
    return {
        call_id: {**result, "call_id": call_id}
        for call_id, result in _sources().items()
    }


def test_generate_report_creates_report_with_valid_sources(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The MCP tool builds a DOCX only after the cited evidence validates."""
    reports_dir = tmp_path / "reports"
    monkeypatch.setattr(business_data_server, "REPORTS_DIR", reports_dir)

    result = business_data_server.generate_report(
        _report().model_dump(mode="json"),
        _sources_by_call_id(),
    )

    assert "error" not in result
    report_path = Path(result["path"])
    assert report_path.parent == reports_dir
    assert report_path.is_file()
    assert report_path.suffix == ".docx"
    assert report_path.name.startswith("monthly-sales-report-november-2024-")
    assert result["warnings"] == []

    document = Document(str(report_path))
    assert len(document.tables) == 3
    assert "R184,200" in document.tables[0].cell(0, 0).text
    assert document.inline_shapes


def test_generate_report_rejects_invalid_object_without_creating_file(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """An invalid model returns a useful error before making output files."""
    reports_dir = tmp_path / "reports"
    monkeypatch.setattr(business_data_server, "REPORTS_DIR", reports_dir)

    result = business_data_server.generate_report({}, _sources_by_call_id())

    assert "error" in result
    assert "Invalid report object" in result["error"]
    assert not reports_dir.exists()


def test_generate_report_rejects_unsupported_evidence_without_creating_file(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Altered claims return the named evidence issue and save no report."""
    reports_dir = tmp_path / "reports"
    monkeypatch.setattr(business_data_server, "REPORTS_DIR", reports_dir)
    report = _report()
    report.key_figures[0].value = 19900000

    result = business_data_server.generate_report(
        report.model_dump(mode="json"),
        _sources_by_call_id(),
    )

    assert "error" in result
    assert "key_figures[0]" in result["error"]
    assert "199,000" in result["error"]
    assert not reports_dir.exists()


def test_generate_report_requires_tool_evidence(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Direct calls without the originating tool results fail closed."""
    reports_dir = tmp_path / "reports"
    monkeypatch.setattr(business_data_server, "REPORTS_DIR", reports_dir)

    result = business_data_server.generate_report(_report().model_dump(mode="json"))

    assert "error" in result
    assert "evidence is unavailable" in result["error"]
    assert not reports_dir.exists()
